# -*- coding: utf-8 -*-
"""
PUBG Mobile chust bazasi — Jurat uchun. Natija o'yindagi "Настройки → Чувствительность" ekrani ko'rinishida (rasm + ruscha yozuvlar).
Ma'lumot: pubg_baza.json (shu papkada). Rasm uchun: pip install pillow

Bo'limlar (bolim):  kamera | ads | giro | giroads
Qatorlar (qator):   3l (3-е лицо без прицела) | 1l (1-е лицо без прицела) | p3 (Прицел 3-е лицо) | p1 (Прицел 1-е лицо)
                    kol (Коллиматор/голо/мушка/боковой) | 2x | 3x | 4x | 6x | 8x

  python pubg_tool.py video   --id v12 --manba "ORG TERROR" [--kod 1-7690-...] [--fps 120] [--giro ha] [--rejim TDM] [--izoh ..]
  python pubg_tool.py sozlama --id v12 --bolim ads --qator 3x --qiymat 31 [--qurol M416]     (videoda ko'ringan sozlama)
  python pubg_tool.py otish   --id v12 --qurol M416 --pritsel 3x|pritselsiz --zajim ha --headshot 4 --tegdi 9 --siljish tepaga [--masofa 50] [--uq 30] [--vaqt 0:12]
  python pubg_tool.py meniki  --bolim ads --qator 3x --qiymat 31 [--qurol M416]               (Juratning o'z sozlamasini yangilash)
  python pubg_tool.py chust   [--qurol M416] [--rasm]     -> tavsiya: o'yin ekranidagi tartib va nomlar bilan + PNG rasm
  python pubg_tool.py rasm    [--qurol M416] [--meniki]   -> faqat rasm (tavsiya yoki o'z sozlamasi)
  python pubg_tool.py hisobot | excel | list | del --id v12
"""
import argparse, json, os, re, sys, collections, datetime

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "pubg_baza.json")

BOLIMLAR = [  # (kalit, sarlavha, slayder maksimumi)
    ("kamera", "Чувствительность камеры", 300),
    ("ads", "Чувствительность прицела", 300),
    ("giro", "Чувствительность гироскопа", 400),
    ("giroads", "Чувствительность гироскопа ADS", 400),
]
QATORLAR = [
    ("3l", "3-е лицо (без прицела)"),
    ("1l", "1-е лицо (без прицела)"),
    ("p3", "Прицел (3-е лицо)"),
    ("p1", "Прицел (1-е лицо)"),
    ("kol", "Коллиматорный, голографический, мушка и боковой прицелы"),
    ("2x", "Оптический прицел 2×"),
    ("3x", "Прицел 3×, Win94"),
    ("4x", "Прицел 4×, ВСС"),
    ("6x", "Оптический прицел 6×"),
    ("8x", "Оптический прицел 8×"),
]
QNOM = dict(QATORLAR); BNOM = {k: n for k, n, _ in BOLIMLAR}; BMAX = {k: m for k, _, m in BOLIMLAR}
ALIAS_Q = {"3lico": "3l", "3-lico": "3l", "tpp": "3l", "1lico": "1l", "1-lico": "1l", "fpp": "1l",
           "pritsel3": "p3", "pritsel1": "p1", "kollimator": "kol", "redot": "kol", "golo": "kol", "1x": "kol"}
ALIAS_B = {"kamera": "kamera", "obyektiv": "kamera", "ads": "ads", "pritsel": "ads", "giro": "giro", "giroskop": "giro",
           "giroads": "giroads", "giroskopads": "giroads"}

# Juratning sozlamasi — 16.10.2026 skrinshotlaridan (M416 shaxsiy konfiguratsiyasi, "Продвинутые настройки")
MENIKI_DEFAULT = {
    "kamera": {"3l": 160, "1l": 140, "p3": 80, "p1": 75, "kol": 75, "2x": 42, "3x": 30, "4x": 23, "6x": 15, "8x": 11},
    "ads": {"3l": 140, "1l": 50, "p3": 70, "p1": 1, "kol": 65, "2x": 17, "3x": 31, "4x": 25, "6x": 16, "8x": 12},
    "giro": {"3l": 390, "1l": 370, "p3": 380, "p1": 350, "kol": 380, "2x": 330, "3x": 270},
    "giroads": {"3l": 395, "1l": 390, "p3": 385, "p1": 370, "kol": 385, "2x": 330, "3x": 260, "4x": 215, "6x": 130},
}


def load():
    if not os.path.exists(DATA):
        return {"videolar": {}, "sozlamalar": [], "otishlar": [], "meniki": json.loads(json.dumps(MENIKI_DEFAULT))}
    with open(DATA, encoding="utf-8") as f: db = json.load(f)
    db.setdefault("meniki", json.loads(json.dumps(MENIKI_DEFAULT))); return db


def save(db):
    tmp = DATA + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f: json.dump(db, f, ensure_ascii=False, indent=1)
    os.replace(tmp, DATA)


def nb(b):
    k = re.sub(r"[^a-z]", "", b.lower())
    if k not in ALIAS_B: raise SystemExit("bolim: kamera | ads | giro | giroads")
    return ALIAS_B[k]


def nqt(q):
    k = q.lower().replace(" ", "").replace("х", "x").replace("×", "x")
    k = ALIAS_Q.get(k, k)
    if k not in QNOM: raise SystemExit("qator: " + " | ".join(QNOM))
    return k


def npr(p):
    k = p.lower().replace(" ", "").replace("х", "x")
    if k in ("pritselsiz", "bez", "bezpritsel", "hip", "hipfire", "yoq", "0x"): return "pritselsiz"
    return nqt(k)


def nq(q): return (q or "").strip().upper().replace(" ", "")


# ---------- commands ----------
def cmd_video(o, db):
    db["videolar"][o.id] = {"manba": o.manba, "kod": o.kod, "fps": o.fps, "giro": o.giro, "rejim": o.rejim, "izoh": o.izoh,
                            "sana": datetime.date.today().strftime("%d.%m.%Y")}
    save(db); print(f"Video {o.id} yozildi ({o.manba}{', kod ' + o.kod if o.kod else ''})")


def cmd_sozlama(o, db):
    r = {"id": o.id, "bolim": nb(o.bolim), "qator": nqt(o.qator), "qiymat": float(o.qiymat), "qurol": nq(o.qurol)}
    db["sozlamalar"] = [x for x in db["sozlamalar"] if not all(x[k] == r[k] for k in ("id", "bolim", "qator", "qurol"))]
    db["sozlamalar"].append(r); save(db)
    print(f"{o.id} | {r['qurol'] or 'umumiy'} | {BNOM[r['bolim']]} → {QNOM[r['qator']]} = {r['qiymat']:g}%")


def cmd_meniki(o, db):
    b, q = nb(o.bolim), nqt(o.qator)
    old = db["meniki"].setdefault(b, {}).get(q)
    db["meniki"][b][q] = float(o.qiymat); save(db)
    print(f"Sizning sozlamangiz: {BNOM[b]} → {QNOM[q]}: {old if old is not None else '-'} → {float(o.qiymat):g}%")


def cmd_otish(o, db):
    r = {"id": o.id, "qurol": nq(o.qurol), "pritsel": npr(o.pritsel), "zajim": o.zajim.lower().startswith("h"),
         "masofa": float(o.masofa) if o.masofa else None, "headshot": int(o.headshot or 0), "tegdi": int(o.tegdi or 0),
         "uq": int(o.uq or 0), "siljish": o.siljish, "vaqt": o.vaqt}
    db["otishlar"].append(r); save(db)
    pn = "без прицела" if r["pritsel"] == "pritselsiz" else QNOM[r["pritsel"]]
    print(f"Otish {o.id} {o.vaqt}: {r['qurol']} | {pn} | {'zajim' if r['zajim'] else 'bir-bir'} | boshga {r['headshot']}" +
          (f"/{r['tegdi']}" if r["tegdi"] else "") + f" | siljish: {r['siljish']}")


def score(db, vid):
    s = [x for x in db["otishlar"] if x["id"] == vid]
    return sum(x["headshot"] for x in s), sum(x["tegdi"] for x in s)


def tavsiya(db, qurol=""):
    """{bolim: {qator: (qiymat, izoh, n_video)}} — videolardagi sozlamalar (headshot bilan og'irlangan) + otish kuzatuvi tuzatishi."""
    q = nq(qurol); out = collections.defaultdict(dict)
    for b, _, mx in BOLIMLAR:
        for k, _ in QATORLAR:
            v = [(s["qiymat"], 1 + score(db, s["id"])[0]) for s in db["sozlamalar"]
                 if s["bolim"] == b and s["qator"] == k and (s["qurol"] in ("", q) or not q)]
            if not v: continue
            rec = sum(a * w for a, w in v) / sum(w for _, w in v); note = ""
            if k in ("3l", "1l") and b in ("kamera", "giro"):
                obs = [x for x in db["otishlar"] if x["pritsel"] == "pritselsiz" and (not q or x["qurol"] == q)]
            elif b in ("ads", "giroads"):
                obs = [x for x in db["otishlar"] if x["pritsel"] == k and (not q or x["qurol"] == q)]
            else:
                obs = []
            up = sum(1 for x in obs if x["siljish"] == "tepaga"); dn = sum(1 for x in obs if x["siljish"] == "pastga")
            if up > dn: rec *= 1.05; note = f"+5%: otishda tepaga ketgan ({up})"
            elif dn > up: rec *= 0.95; note = f"−5%: pastga tushgan ({dn})"
            out[b][k] = (min(round(rec), mx), note, len(v))
    return out


def cmd_chust(o, db):
    t = tavsiya(db, o.qurol); mine = db["meniki"]
    if not t: print("Baza bo'sh — avval video tahlil qiling (sozlama ekrani ko'ringan video kerak)."); return
    print(f"ЧУВСТВИТЕЛЬНОСТЬ{' — ' + nq(o.qurol) if o.qurol else ''} (tavsiya | sizda | farq)")
    for b, title, _ in BOLIMLAR:
        if b not in t: continue
        print(f"\n▌{title}")
        for k, nom in QATORLAR:
            if k not in t[b]: continue
            val, note, n = t[b][k]; m = mine.get(b, {}).get(k)
            d = "" if m is None else (f" | sizda {m:g}% → {val - m:+g}" if val != m else " | sizda ham shu ✓")
            print(f"  {nom} — {val}%{d}{'  (' + note + ')' if note else ''}")
    hip = [x for x in db["otishlar"] if x["pritsel"] == "pritselsiz" and (not o.qurol or x["qurol"] == nq(o.qurol))]
    if hip:
        hs = sum(x["headshot"] for x in hip); tg = sum(x["tegdi"] for x in hip)
        print(f"\nБЕЗ ПРИЦЕЛА: {len(hip)} epizod, boshga {hs}" + (f"/{tg} ({hs * 100 // max(tg, 1)}%)" if tg else "") +
              " — Чувствительность камеры va гироскопа: 3-е/1-е лицо (без прицела) qatorlari")
    yoq = [QNOM[k] for k, _ in QATORLAR if not any(k in t[b] for b in t)]
    if yoq: print("\nMa'lumot yo'q: " + "; ".join(yoq))
    if o.rasm: print(rasm(db, toliq(db, t), o.qurol, "Рекомендуемая чувствительность"))


def cmd_rasm(o, db):
    if o.meniki:
        t = {b: {k: (v, "", 0) for k, v in d.items()} for b, d in db["meniki"].items()}
        print(rasm(db, t, o.qurol, "Моя чувствительность", compare=False))
    else:
        print(rasm(db, toliq(db, tavsiya(db, o.qurol)), o.qurol, "Рекомендуемая чувствительность"))


def toliq(db, t):
    """Tavsiya bo'lmagan qatorlarga Juratning o'z qiymatini qo'yadi — rasmda butun ekran ko'rinadi."""
    out = {b: {k: (v, "", 0) for k, v in d.items()} for b, d in db["meniki"].items()}
    for b, d in t.items():
        for k, v in d.items(): out.setdefault(b, {})[k] = v
    return out


# ---------- o'yin ekrani ko'rinishidagi rasm ----------
def font(size, bold=True):
    from PIL import ImageFont
    cands = ([r"C:\Windows\Fonts\bahnschrift.ttf", r"C:\Windows\Fonts\ARIALNB.TTF", r"C:\Windows\Fonts\arialbd.ttf"] +
             ["/usr/share/fonts/truetype/dejavu/DejaVuSansCondensed-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSansCondensed.ttf"])
    for c in cands:
        if os.path.exists(c):
            try:
                f = ImageFont.truetype(c, size)
                if c.lower().endswith("bahnschrift.ttf"):
                    try: f.set_variation_by_name("SemiBold Condensed" if bold else "Condensed")
                    except Exception: pass
                return f
            except Exception: continue
    return ImageFont.load_default()


def rasm(db, t, qurol, title, compare=True):
    from PIL import Image, ImageDraw
    W, ROW, HEAD = 1700, 100, 78
    rows = [(b, k) for b, _, _ in BOLIMLAR for k, _ in QATORLAR if k in t.get(b, {})]
    secs = [b for b, _, _ in BOLIMLAR if t.get(b)]
    H = 202 + len(secs) * (HEAD + 20) + len(rows) * ROW + 70
    im = Image.new("RGB", (W, H), (38, 38, 40)); d = ImageDraw.Draw(im)
    fT, fH, fR, fV, fS = font(40), font(34), font(32), font(34), font(24, False)
    WHITE, GREY, ROWBG, YEL = (235, 235, 235), (150, 150, 150), (52, 52, 54), (240, 180, 40)
    # tepa: tab va sarlavha
    ft = font(28); tw0 = d.textlength("ПРОДВИНУТЫЕ НАСТРОЙКИ", font=ft)
    d.rectangle([40, 30, 80 + tw0, 92], fill=(200, 200, 200)); d.text((60, 44), "ПРОДВИНУТЫЕ НАСТРОЙКИ", font=ft, fill=(20, 20, 20))
    d.line([40, 92, 80 + tw0, 92], fill=(60, 80, 255), width=4)
    hdr = title + (f" для {nq(qurol)}" if qurol else "")
    d.rectangle([40, 112, W - 40, 112 + 70], fill=(56, 56, 58)); d.rectangle([40, 112, 46, 182], fill=YEL)
    d.text((70, 128), hdr, font=fT, fill=WHITE)
    y = 112 + 70 + 20
    mine = db.get("meniki", {})
    for b, btitle, mx in BOLIMLAR:
        if not t.get(b): continue
        d.rectangle([40, y, W - 40, y + HEAD], fill=(56, 56, 58)); d.rectangle([40, y, 46, y + HEAD], fill=YEL)
        d.text((70, y + 20), btitle, font=fH, fill=WHITE); y += HEAD + 6
        for k, nom in QATORLAR:
            if k not in t[b]: continue
            val = t[b][k][0]
            d.rectangle([40, y, W - 40, y + ROW - 6], fill=ROWBG)
            fl, sz = fR, 32
            while d.textlength(nom, font=fl) > 800 and sz > 20: sz -= 2; fl = font(sz)
            d.text((70, y + 30 + (32 - sz) // 2), nom, font=fl, fill=WHITE)
            m = mine.get(b, {}).get(k)
            if compare and m is not None and m != val:
                dt = f"{val - m:+g}"; d.text((885, y + 36), f"было {m:g}%", font=fS, fill=GREY)
                d.text((1020, y + 36), dt, font=fS, fill=(90, 210, 120) if val > m else (240, 110, 90))
            vs = f"{val:g}%"; tw = d.textlength(vs, font=fV); d.text((1185 - tw, y + 28), vs, font=fV, fill=WHITE)
            # [-] slider [+]
            d.rectangle([1205, y + 22, 1255, y + 72], outline=(120, 120, 120), width=2); d.line([1218, y + 47, 1242, y + 47], fill=WHITE, width=4)
            x0, x1 = 1280, 1560; cy = y + 47
            d.rectangle([x0, cy - 8, x1, cy + 8], fill=(80, 80, 82))
            px = x0 + (x1 - x0) * max(0, min(val, mx)) / mx
            d.rectangle([x0, cy - 8, px, cy + 8], fill=(150, 150, 152)); d.rectangle([px - 8, cy - 18, px + 8, cy + 18], fill=(120, 120, 122))
            d.rectangle([1585, y + 22, 1635, y + 72], outline=(120, 120, 120), width=2)
            d.line([1598, y + 47, 1622, y + 47], fill=WHITE, width=4); d.line([1610, y + 35, 1610, y + 59], fill=WHITE, width=4)
            y += ROW
        y += 14
    d.text((40, H - 45), f"Jurat Assistant · {datetime.date.today():%d.%m.%Y}", font=fS, fill=GREY)
    name = f"chust_{nq(qurol) or 'umumiy'}_{'meniki' if not compare else 'tavsiya'}_{datetime.datetime.now():%Y%m%d_%H%M}.png"
    p = os.path.join(HERE, name); im.save(p); return p


def cmd_hisobot(o, db):
    if not db["videolar"]: print("Video yo'q."); return
    for vid, v in db["videolar"].items():
        s = [x for x in db["otishlar"] if x["id"] == vid]; z = [x for x in db["sozlamalar"] if x["id"] == vid]
        hs, tg = score(db, vid); hip = [x for x in s if x["pritsel"] == "pritselsiz"]
        print(f"{vid} | {v['manba']} | kod: {v.get('kod') or '-'} | FPS {v.get('fps') or '?'} | гироскоп {v.get('giro') or '?'}")
        print(f"   qurollar: {', '.join(sorted({x['qurol'] for x in s})) or '-'} | zajim {sum(x['zajim'] for x in s)} | headshot {hs}" + (f"/{tg}" if tg else ""))
        if hip: print(f"   БЕЗ ПРИЦЕЛА: {len(hip)} epizod, boshga {sum(x['headshot'] for x in hip)}")
        if z: print(f"   sozlama ekrani: {len(z)} qiymat")


def cmd_excel(o, db):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    F = Font(name="Arial", size=10); H = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    HF = PatternFill("solid", fgColor="2B2B2D"); YF = PatternFill("solid", fgColor="F0B428")
    th = Side(style="thin", color="BFBFBF"); BD = Border(left=th, right=th, top=th, bottom=th)
    wb = Workbook(); ws = wb.active; ws.title = "Чувствительность"; t = tavsiya(db); mine = db["meniki"]
    for i, c in enumerate(["Раздел", "Параметр", "Моя", "Рекомендация", "Разница"], 1):
        x = ws.cell(1, i, c); x.font = H; x.fill = HF; x.border = BD; x.alignment = Alignment(horizontal="center")
    r = 2
    for b, title, _ in BOLIMLAR:
        x = ws.cell(r, 1, title); x.font = Font(name="Arial", size=10, bold=True); x.fill = YF; r += 1
        for k, nom in QATORLAR:
            m = mine.get(b, {}).get(k); rec = t.get(b, {}).get(k, (None,))[0]
            if m is None and rec is None: continue
            for c, v in enumerate([title, nom, m, rec, (rec - m) if (m is not None and rec is not None) else None], 1):
                x = ws.cell(r, c, v); x.font = F; x.border = BD
            r += 1
    for c, w in zip("ABCDE", [30, 52, 10, 14, 10]): ws.column_dimensions[c].width = w
    o2 = wb.create_sheet("Otishlar")
    for i, c in enumerate(["Video", "Vaqt", "Qurol", "Pritsel", "Zajim", "Masofa", "Tegdi", "Headshot", "Siljish"], 1): o2.cell(1, i, c).font = Font(bold=True)
    for j, x in enumerate(db["otishlar"], 2):
        for i, v in enumerate([x["id"], x["vaqt"], x["qurol"], "без прицела" if x["pritsel"] == "pritselsiz" else QNOM[x["pritsel"]],
                               "ha" if x["zajim"] else "yo'q", x["masofa"], x["tegdi"], x["headshot"], x["siljish"]], 1): o2.cell(j, i, v)
    p = os.path.join(HERE, "PUBG_chust.xlsx"); wb.save(p); print(p)


def main():
    ap = argparse.ArgumentParser(); sp = ap.add_subparsers(dest="cmd", required=True)
    v = sp.add_parser("video"); v.add_argument("--id", required=True); v.add_argument("--manba", required=True)
    for n in ("--kod", "--fps", "--giro", "--rejim", "--izoh"): v.add_argument(n, default="")
    for name in ("sozlama", "meniki"):
        s = sp.add_parser(name)
        if name == "sozlama": s.add_argument("--id", required=True)
        for n in ("--bolim", "--qator", "--qiymat"): s.add_argument(n, required=True)
        s.add_argument("--qurol", default="")
    t = sp.add_parser("otish")
    for n in ("--id", "--qurol", "--pritsel"): t.add_argument(n, required=True)
    for n, dflt in (("--zajim", "yoq"), ("--masofa", ""), ("--headshot", "0"), ("--tegdi", "0"), ("--uq", "0"), ("--siljish", "barqaror"), ("--vaqt", "")):
        t.add_argument(n, default=dflt)
    c = sp.add_parser("chust"); c.add_argument("--qurol", default=""); c.add_argument("--rasm", action="store_true")
    r = sp.add_parser("rasm"); r.add_argument("--qurol", default=""); r.add_argument("--meniki", action="store_true")
    sp.add_parser("hisobot"); sp.add_parser("excel"); sp.add_parser("list")
    dl = sp.add_parser("del"); dl.add_argument("--id", required=True)
    o = ap.parse_args(); db = load()
    if o.cmd == "list": print(json.dumps(db, ensure_ascii=False, indent=1)); return
    if o.cmd == "del":
        db["videolar"].pop(o.id, None)
        for k in ("sozlamalar", "otishlar"): db[k] = [x for x in db[k] if x["id"] != o.id]
        save(db); print("o'chirildi"); return
    {"video": cmd_video, "sozlama": cmd_sozlama, "meniki": cmd_meniki, "otish": cmd_otish, "chust": cmd_chust,
     "rasm": cmd_rasm, "hisobot": cmd_hisobot, "excel": cmd_excel}[o.cmd](o, db)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"): sys.stdout.reconfigure(encoding="utf-8")
    main()
