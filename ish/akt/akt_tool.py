# -*- coding: utf-8 -*-
"""
AKT SVERKA dasturi — Jurat uchun. Bank ko'chirmalari va fakturalar bir marta yuklanadi (baza), keyin
istalgan kontragent / davr / shartnoma uchun klassik "АКТ СВЕРКИ ВЗАИМНЫХ РАСЧЕТОВ" Excel yasaladi.

Baza fayllari (shu papkada): akt_baza.json, firmalar.json (bizning kompaniyalar rekvizitlari).

  python akt_tool.py bank    --biz TSB --file vypiska.html      (Xalq banki HTML/CP1251, yoki .xlsx/.csv) — takrorlar o'tkazib yuboriladi
  python akt_tool.py faktura --biz TSB --file didox.xlsx        (Didox / Soliq servis eksporti; .xls uchun: pip install xlrd)
  python akt_tool.py qator   --biz TSB --stir 123456789 --date 05.03.2026 --tur sotdik|oldik|toladik|toladi --summa 15000000 --hujjat "Faktura 45" [--shartnoma ..] [--nom ..]
        sotdik  = biz tovar/xizmat berdik (bizning faktura)      -> bizda DEBET
        oldik   = ular bizga berdi (ularning fakturasi)          -> bizda KREDIT
        toladik = biz ularga pul o'tkazdik                       -> bizda DEBET
        toladi  = ular bizga pul o'tkazdi                        -> bizda KREDIT
  python akt_tool.py kontragentlar --biz TSB [--qidir bo'ston]   -> bazadagi kontragentlar (STIR, nom, qatorlar soni)
  python akt_tool.py akt --biz TSB --stir 123456789 --dan 01.01.2026 --gacha 30.09.2026 [--saldo 0] [--shartnoma 22-ГБ-2026] [--nom "..."] [--sana 30.09.2026]
        --saldo: davr boshidagi saldo (+ = kontragent bizga qarz, − = biz qarzmiz). Berilmasa — bazadagi oldingi qatorlardan hisoblanadi.
  python akt_tool.py tashlangan --biz TSB [--stir ..]            -> fakturadan nima tashlangani va nega
  python akt_tool.py del --id N | holat | tozala --biz TSB --manba bank   (noto'g'ri yuklangan bank qatorlarini o'chirish)
"""
import argparse, json, os, re, sys, datetime, collections, csv, html
from html.parser import HTMLParser

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "akt_baza.json")
FIRMS = os.path.join(HERE, "firmalar.json")
DEFAULT_FIRMS = {
    "TSB": {"nom": "ООО \"TASHKENT STROY BUILD\"", "stir": "307669546", "direktor": "Qazoqov T.R.", "buxgalter": "",
            "eski_nom": "TEMUR TEXNIK TRANS"},
    "TLT": {"nom": "ООО \"TOSHKENT LOGISTIK TRANS\"", "stir": "310152233", "direktor": "Ravshanov J.", "buxgalter": ""},
    "RAYXONA": {"nom": "ООО \"RAYXONA SHIRINA SERVIS\"", "stir": "", "direktor": "", "buxgalter": ""},
}


# ---------------- util ----------------
def firms():
    if not os.path.exists(FIRMS):
        with open(FIRMS, "w", encoding="utf-8") as f: json.dump(DEFAULT_FIRMS, f, ensure_ascii=False, indent=1)
    with open(FIRMS, encoding="utf-8") as f: return json.load(f)


def load():
    if not os.path.exists(DATA): return {"qatorlar": [], "tashlangan": [], "nomlar": {}}
    with open(DATA, encoding="utf-8") as f: return json.load(f)


def save(db):
    tmp = DATA + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f: json.dump(db, f, ensure_ascii=False, indent=1)
    os.replace(tmp, DATA)


def to_iso(s):
    if isinstance(s, (datetime.date, datetime.datetime)): return s.strftime("%Y-%m-%d")
    s = str(s).strip()
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s)          # 2021-01-04 15:25:00 (Excel sanasi)
    if m: return "-".join(m.groups())
    m = re.search(r"(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})", s)
    if m:
        d, mo, y = m.groups(); y = ("20" + y) if len(y) == 2 else y
        return f"{y}-{int(mo):02d}-{int(d):02d}"
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m: return "-".join(m.groups())
    return None


def dmy(iso): return f"{iso[8:10]}.{iso[5:7]}.{iso[:4]}"


def num(x):
    if x is None: return None
    if isinstance(x, (int, float)): return float(x)
    s = str(x).replace(" ", "").replace(" ", "").replace("'", "")
    if not s or s in ("-", "—"): return None
    if s.count(",") == 1 and s.count(".") == 0: s = s.replace(",", ".")
    elif s.count(",") >= 1 and s.count(".") == 1: s = s.replace(",", "")
    try: return float(s)
    except ValueError: return None


def stirs(text):
    return re.findall(r"(?<!\d)(\d{9})(?!\d)", str(text or ""))


def idn(x):
    """INN (9 raqam) yoki PINFL (14 raqam) — katakda faqat raqam bo'lsa."""
    d = re.sub(r"\D", "", str(x or ""))
    if len(d) in (9, 14): return d
    f = stirs(x); return f[0] if f else ""


def money(x): return f"{x:,.2f}".replace(",", " ")


def add(db, r):
    if r.get("fid"):
        if any(x.get("fid") == r["fid"] and x["biz"] == r["biz"] for x in db["qatorlar"]): return False
    else:
        key = (r["biz"], r["stir"], r["date"], r["tur"], round(r["summa"], 2), r.get("hujjat", ""))
        for x in db["qatorlar"]:
            if not x.get("fid") and (x["biz"], x["stir"], x["date"], x["tur"], round(x["summa"], 2), x.get("hujjat", "")) == key: return False
    r["id"] = max([x["id"] for x in db["qatorlar"]] or [0]) + 1
    db["qatorlar"].append(r); return True


# ---------------- reading tables ----------------
class TP(HTMLParser):
    def __init__(self):
        super().__init__(); self.rows = []; self.row = None; self.cell = None
    def handle_starttag(self, t, a):
        if t == "tr": self.row = []
        elif t in ("td", "th") and self.row is not None: self.cell = []
        elif t == "br" and self.cell is not None: self.cell.append(" ")
    def handle_endtag(self, t):
        if t in ("td", "th") and self.cell is not None and self.row is not None:
            self.row.append(re.sub(r"\s+", " ", "".join(self.cell)).strip()); self.cell = None
        elif t == "tr" and self.row is not None:
            if any(self.row): self.rows.append(self.row)
            self.row = None
    def handle_data(self, d):
        if self.cell is not None: self.cell.append(d)


def read_rows(path):
    ext = os.path.splitext(path)[1].lower()
    head = open(path, "rb").read(2048).lstrip().lower()
    if ext in (".xls", ".xlsx", ".csv") and (head.startswith(b"<") or b"<html" in head or b"<table" in head):
        ext = ".html"                                   # bank "xls" deb beradi, aslida HTML
    if ext in (".html", ".htm"):
        raw = open(path, "rb").read()
        for enc in ("cp1251", "utf-8"):
            try: txt = raw.decode(enc); break
            except UnicodeDecodeError: continue
        if "charset=utf-8" in txt[:2000].lower():
            txt = raw.decode("utf-8", errors="replace")
        p = TP(); p.feed(html.unescape(txt)); return p.rows
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True); out = []
        for ws in wb.worksheets:
            try: ws.reset_dimensions()          # ba'zi eksportlarda o'lcham noto'g'ri yozilgan (faqat 1 qator ko'rinadi)
            except Exception: pass
            for r in ws.iter_rows(values_only=True):
                if any(v not in (None, "") for v in r): out.append(list(r))
        return out
    if ext == ".xls":
        try: import xlrd
        except ImportError: raise SystemExit(".xls uchun: python -m pip install xlrd  (yoki faylni .xlsx qilib saqlang)")
        wb = xlrd.open_workbook(path); out = []
        for sh in wb.sheets():
            for i in range(sh.nrows):
                r = []
                for c in sh.row(i):
                    v = c.value
                    if c.ctype == 3: v = datetime.datetime(*xlrd.xldate_as_tuple(v, wb.datemode))
                    r.append(v)
                if any(v not in (None, "") for v in r): out.append(r)
        return out
    if ext == ".csv":
        raw = open(path, "rb").read()
        try: txt = raw.decode("utf-8-sig")
        except UnicodeDecodeError: txt = raw.decode("cp1251")
        return [r for r in csv.reader(txt.splitlines(), delimiter=";" if txt.count(";") > txt.count(",") else ",")]
    raise SystemExit("Fayl turi: .html .xlsx .xls .csv")


def find_header(rows, must):
    for i, r in enumerate(rows[:60]):
        low = [str(c or "").lower().replace("ё", "е") for c in r]
        if all(any(k in c for c in low) for k in must): return i, low
    return None, None


def col(low, *keys, exclude=()):
    for k in keys:
        for j, c in enumerate(low):
            if k in c and not any(e in c for e in exclude): return j
    return None


# ---------------- import: bank ----------------
def cmd_bank(o, db):
    F = firms()[o.biz]; our = F["stir"]
    rows = read_rows(o.file)
    hi, low = find_header(rows, ["дебет", "кредит"])
    if hi is None: hi, low = find_header(rows, ["debet", "kredit"])
    if hi is not None:
        cD = col(low, "дебет", "debet"); cK = col(low, "кредит", "kredit")
        cDate = col(low, "дата", "sana"); cDoc = col(low, "№", "номер", "hujjat")
        cName = col(low, "наимен", "корресп", "контраг", "получат", "платель", "nomi")
        cInn = col(low, "инн", "stir"); cPur = col(low, "назнач", "maqsad", "детали")
    else:
        # Xalq banki sarlavhasiz formati: Дата | Счет/ИНН/Наименование | № док | Оп | МФО | Дебет | Кредит | Назначение
        ok = [r for r in rows if len(r) >= 8 and to_iso(r[0]) and "/" in str(r[1] or "")]
        if len(ok) < max(3, len(rows) // 3):
            raise SystemExit("Bank ko'chirmasida 'Дебет/Кредит' sarlavhasi topilmadi va format tanilmadi — faylni Juratdan so'rang")
        hi = -1; cDate, cInn, cDoc, cD, cK, cPur, cName = 0, 1, 2, 5, 6, 7, None
    add_n = dup = skip = own = fixed = 0; cur_date = None; sD = sK = 0.0
    # INN'i "000000000" yoki bo'sh kelgan qatorlar uchun: hisob raqami / nomi bo'yicha haqiqiy STIRni topish
    acc2stir = dict(db.get("hisoblar", {})); nom2stir = {}
    for st_, nm_ in db["nomlar"].items(): nom2stir[re.sub(r"[^A-ZА-Я0-9]", "", nm_.upper().replace("MCHJ", "").replace("МЧЖ", "").replace("OOO", "").replace("ООО", ""))] = st_
    if cInn is not None:
        for r0 in rows[hi + 1:]:
            pp = [x.strip() for x in str(r0[cInn] if cInn < len(r0) else "").split("/")]
            if len(pp) >= 2:
                dd = re.sub(r"\D", "", pp[1])
                if len(dd) in (9, 14) and dd.strip("0"):
                    acc2stir.setdefault(pp[0], dd)
                    if len(pp) >= 3: nom2stir.setdefault(re.sub(r"[^A-ZА-Я0-9]", "", pp[2].upper().replace("MCHJ", "").replace("МЧЖ", "").replace("OOO", "").replace("ООО", "")), dd)
    db["hisoblar"] = acc2stir
    for r in rows[hi + 1:]:
        raw = list(r)
        cells = [str(c if c is not None else "") for c in r]
        g = lambda j: cells[j] if (j is not None and j < len(cells)) else ""
        rd = raw[cDate] if (cDate is not None and cDate < len(raw)) else None
        d0 = to_iso(rd) if rd not in (None, "") else None
        d = d0 or cur_date
        if d0: cur_date = d0
        deb, kre = num(g(cD)), num(g(cK))
        if not d or (not deb and not kre):
            skip += 1; continue
        txt = " ".join(cells)
        if not d0 and re.search(r"итого|оборот|сальдо|остаток|jami", txt.lower()): skip += 1; continue
        parts = [x.strip() for x in g(cInn).split("/")]       # Xalq banki: "hisob/INN/nomi"
        if len(parts) >= 2 and re.sub(r"\D", "", parts[1]) and len(re.sub(r"\D", "", parts[1])) in (9, 14):
            cand = [re.sub(r"\D", "", parts[1])]
            if not cand[0].strip("0"):                       # INN 000000000 — hisob/nom bo'yicha tiklash
                nk = re.sub(r"[^A-ZА-Я0-9]", "", (parts[2] if len(parts) >= 3 else "").upper().replace("MCHJ", "").replace("МЧЖ", "").replace("OOO", "").replace("ООО", ""))
                real = acc2stir.get(parts[0]) or (nom2stir.get(nk) if nk else None)
                if real: cand = [real]; fixed += 1
        else:
            cand = stirs(g(cInn)) + stirs(g(cName)) + stirs(txt)
        if cand and all(x == our for x in cand): own += 1; continue
        cand = [x for x in cand if x != our]
        if not cand: skip += 1; continue
        st = cand[0]; nom = g(cName) or (parts[2] if len(parts) >= 3 else "")
        if nom: db["nomlar"].setdefault(st, re.sub(r"\s*\d{9,}.*", "", nom)[:120])
        pur = g(cPur)
        sh = re.search(r"(?:дог\w*|shartnoma|contract)\s*№?\s*([\w\-/\.А-Яа-я]+)", pur, re.I)
        r_ = {"biz": o.biz, "stir": st, "hisob": parts[0] if len(parts) >= 2 else "", "nom": (parts[2] if len(parts) >= 3 else nom)[:120],
              "date": d, "tur": "toladik" if deb else "toladi", "summa": deb or kre,
              "hujjat": ("п/п " + (re.sub(r"\.0$", "", g(cDoc).strip()).lstrip("0") or "0")).strip(), "izoh": pur[:200], "shartnoma": sh.group(1) if sh else "", "manba": "bank"}
        sD += deb or 0; sK += kre or 0
        if add(db, r_): add_n += 1
        else: dup += 1
    save(db)
    print(f"Bank ({o.biz}): {add_n} ta qo'shildi, {dup} ta takror | o'z hisob/komissiya qatorlari: {own} | sarlavha/jami: {skip}" + (f" | INN'i 000000000 bo'lib, hisob/nom bo'yicha tiklangan: {fixed}" if fixed else ""))
    print(f"Kontragentlar bo'yicha: Дебет (biz to'ladik) {money(sD)} | Кредит (bizga tushdi) {money(sK)}")


# ---------------- import: faktura ----------------
def cmd_faktura(o, db):
    F = firms()[o.biz]; our = F["stir"]
    rows = read_rows(o.file)
    hi, low = find_header(rows, ["статус"])
    if hi is None: hi, low = find_header(rows, ["holat"])
    if hi is None: raise SystemExit("Faktura faylida 'Статус' ustuni topilmadi")
    cSt = col(low, "статус", "holat"); cNo = col(low, "№ док", "номер док", "№ сч", "номер", "№", exclude=("договор", "исход"))
    cDate = col(low, "дата док", "дата сч", "дата", exclude=("договор",))
    cSel = col(low, "продавец (инн", "инн продав", "инн поставщ", "продавец", "поставщик", "sotuvchi")
    cBuy = col(low, "покупатель (инн", "инн покуп", "покупатель", "xaridor")
    cSelN = col(low, "продавец (наимен", "наименование продав"); cBuyN = col(low, "покупатель (наимен", "наименование покуп")
    cSum = col(low, "с учетом ндс", "с ндс", "итого", "jami")
    if cSum is None and len(low) >= 36: cSum = 35          # Didox: 36-ustun = QQS bilan jami
    cId = next((j for j, c in enumerate(low) if c.strip() == "id"), None)
    cType = col(low, "тип", "turi"); cCon = col(low, "договор", "shartnoma"); cOld = col(low, "исход", "старый")
    items = []
    for r in rows[hi + 1:]:
        cells = [r[j] if j < len(r) else "" for j in range(len(low))]
        g = lambda j: (cells[j] if j is not None and j < len(cells) else "")
        st = str(g(cSt) or "").strip()
        if not st: continue
        items.append({"st": st, "no": str(g(cNo) or "").strip(), "date": to_iso(g(cDate)), "sel": idn(g(cSel)),
                      "buy": idn(g(cBuy)), "sum": num(g(cSum)), "type": str(g(cType) or ""),
                      "fid": str(g(cId) or "").strip(), "seln": str(g(cSelN) or "").strip(), "buyn": str(g(cBuyN) or "").strip(),
                      "con": str(g(cCon) or "").strip(), "old": str(g(cOld) or "").strip()})
    # Faqat faktura sarlavha qatorlari olinadi (СТАТУС bor); tovar qatorlari o'qilmaydi
    signed = {(i["no"], i["sel"], i["buy"]) for i in items if "подписан" in i["st"].lower()}
    for i in items:
        if i["sel"] and i["seln"]: db["nomlar"].setdefault(i["sel"], i["seln"][:120])
        if i["buy"] and i["buyn"]: db["nomlar"].setdefault(i["buy"], i["buyn"][:120])
    replaced = {i["old"] for i in items if i["old"] and "подписан" in i["st"].lower() and "исправ" in i["type"].lower()}
    add_n = dup = 0; drop = []
    for i in items:
        s = i["st"].lower(); why = ""
        if "подписан" not in s:
            if "не настоящ" in s:
                why = "Не настоящий: tuzatuvchi faktura bilan almashtirilgan (Подписан juftligi hisoblandi)" if (i["no"], i["sel"], i["buy"]) in signed else "holat: Не настоящий"
            else: why = f"holat: {i['st']}"
        elif "без оплат" in i["type"].lower(): why = "Без оплаты (to'lovsiz)"
        elif i["sum"] is None: why = "SUMMA YOZILMAGAN (eksportda yo'q) — summani Juratdan so'rang"
        elif not i["sum"]: why = "nol summa"
        elif i["no"] in replaced: why = "tuzatuvchi faktura bilan almashtirilgan"
        elif our not in (i["sel"], i["buy"]): why = f"bizning STIR ({our}) yo'q"
        if why:
            drop.append({"biz": o.biz, "no": i["no"], "date": i["date"], "sel": i["sel"], "buy": i["buy"], "sum": i["sum"], "sabab": why}); continue
        if not i["date"]: drop.append({"biz": o.biz, "no": i["no"], "sum": i["sum"], "sabab": "sana o'qilmadi"}); continue
        sotdik = i["sel"] == our; other = i["buy"] if sotdik else i["sel"]
        r_ = {"biz": o.biz, "stir": other, "date": i["date"], "tur": "sotdik" if sotdik else "oldik", "summa": i["sum"],
              "hujjat": f"СФ №{i['no']}", "shartnoma": i["con"], "manba": "faktura", "fid": i["fid"], "izoh": i["type"][:60]}
        if add(db, r_): add_n += 1
        else: dup += 1
    db["tashlangan"] = [x for x in db["tashlangan"] if not (x["biz"] == o.biz and any(x.get("no") == d["no"] and x.get("sel") == d.get("sel") for d in drop))] + drop
    save(db)
    cnt = collections.Counter(d["sabab"].split(":")[0].split(" (")[0].split(" —")[0] for d in drop)
    print(f"Faktura ({o.biz}): {add_n} ta qo'shildi, {dup} ta takror, {len(drop)} ta tashlandi" + (": " + ", ".join(f"{k} {v}" for k, v in cnt.items()) if drop else ""))
    print(f"Tekshiruv: fayldagi {len(items)} = qo'shildi {add_n} + takror {dup} + tashlandi {len(drop)}" + (" ✓" if len(items) == add_n + dup + len(drop) else " ✗ — FARQ BOR"))


def cmd_qator(o, db):
    if o.tur not in ("sotdik", "oldik", "toladik", "toladi"): raise SystemExit("--tur: sotdik | oldik | toladik | toladi")
    if o.nom: db["nomlar"][o.stir] = o.nom
    r = {"biz": o.biz, "stir": o.stir, "date": to_iso(o.date), "tur": o.tur, "summa": num(o.summa), "hujjat": o.hujjat,
         "shartnoma": o.shartnoma, "manba": "qo'lda", "izoh": ""}
    print("Qo'shildi" if add(db, r) else "TAKROR — qo'shilmadi"); save(db)


def cmd_kontragentlar(o, db):
    c = collections.Counter(x["stir"] for x in db["qatorlar"] if x["biz"] == o.biz)
    for st, n in c.most_common():
        nm = db["nomlar"].get(st, "")
        if o.qidir and o.qidir.lower() not in (nm + st).lower(): continue
        print(f"{st} | {nm} | {n} qator")


# ---------------- Russian amount in words ----------------
def propis(n):
    cents = int(round(abs(n) * 100)); n, tiy = cents // 100, cents % 100
    if n == 0: return f"Ноль сум {tiy:02d} тийин"
    ones = ["", "один", "два", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"]
    onesf = ["", "одна", "две"] + ones[3:]
    teens = ["десять", "одиннадцать", "двенадцать", "тринадцать", "четырнадцать", "пятнадцать", "шестнадцать", "семнадцать", "восемнадцать", "девятнадцать"]
    tens = ["", "", "двадцать", "тридцать", "сорок", "пятьдесят", "шестьдесят", "семьдесят", "восемьдесят", "девяносто"]
    hund = ["", "сто", "двести", "триста", "четыреста", "пятьсот", "шестьсот", "семьсот", "восемьсот", "девятьсот"]
    groups = [(("", "", ""), False), (("тысяча", "тысячи", "тысяч"), True), (("миллион", "миллиона", "миллионов"), False),
              (("миллиард", "миллиарда", "миллиардов"), False)]

    def tri(x, fem):
        w = [hund[x // 100]]; t = x % 100
        if 10 <= t < 20: w.append(teens[t - 10])
        else: w += [tens[t // 10], (onesf if fem else ones)[t % 10]]
        return [s for s in w if s]

    def form(x, f):
        if 11 <= x % 100 <= 14: return f[2]
        return f[0] if x % 10 == 1 else f[1] if 2 <= x % 10 <= 4 else f[2]
    out = []; gi = 0
    while n > 0 and gi < len(groups):
        x = n % 1000; f, fem = groups[gi]
        if x: out = tri(x, fem) + ([form(x, f)] if f[0] else []) + out
        n //= 1000; gi += 1
    s = " ".join(out) + f" сум {tiy:02d} тийин"
    return s[0].upper() + s[1:]


# ---------------- shartnoma bo'yicha ajratish ----------------
LAT2CYR = str.maketrans({"A": "А", "P": "Р", "C": "С", "Y": "У", "B": "В", "E": "Е", "K": "К", "M": "М", "H": "Н", "O": "О", "T": "Т", "X": "Х"})


def shartnoma_key(x):
    t = (x.get("shartnoma") or "").strip()
    if not t:
        m = re.search(r"(?:№|N)\s*№?\s*([A-Za-zА-Яа-я]{0,4}[\s\-]?[0-9][0-9/\-\.]*)\s*(?:сонли\s*)?(?:шартн|договор|дог\.)", x.get("izoh", ""), re.I) \
            or re.search(r"(?:договор\w*|шартнома\w*)\s*№\s*([A-Za-zА-Яа-я]{0,4}[\s\-]?[0-9][0-9/\-\.]*)", x.get("izoh", ""), re.I) \
            or re.search(r"([A-Za-zА-Яа-я]{1,4}[\s\-]?[0-9][0-9/\-]*)\s*сонли\s*шартн", x.get("izoh", ""), re.I)
        if m: t = m.group(1)
    if not t: return "raqamsiz"
    k = re.sub(r"[\s\-\.№]", "", t.upper()).translate(LAT2CYR)
    return k or "raqamsiz"


def shartnoma_jadval(rows):
    agg = collections.OrderedDict()
    for x in sorted(rows, key=lambda x: x["date"]):
        k = shartnoma_key(x); a = agg.setdefault(k, [0.0, 0.0, x["date"], x["date"]])
        a[0 if x["tur"] in DEBET else 1] += x["summa"]; a[3] = x["date"]
    return agg


# ---------------- AKT ----------------
DEBET = ("sotdik", "toladik")   # bizda Debet (kontragent qarzi oshadi)


def cmd_akt(o, db):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.worksheet.page import PageMargins
    F = firms(); B = F[o.biz]
    _all = [x["date"] for x in db["qatorlar"] if x["biz"] == o.biz and x["stir"] == o.stir]
    # Davr berilmasa: kontragent bilan BIRINCHI operatsiyadan bugungacha (to'liq tarix, saldo 0 dan)
    dan = to_iso(o.dan) if o.dan else (min(_all) if _all else datetime.date.today().strftime("%Y-%m-%d"))
    gacha = to_iso(o.gacha) if o.gacha else datetime.date.today().strftime("%Y-%m-%d")
    if not o.dan and o.saldo in (None, ""): o.saldo = "0"
    allr = [x for x in db["qatorlar"] if x["biz"] == o.biz and x["stir"] == o.stir
            and (not o.shartnoma or o.shartnoma.lower() in (x.get("shartnoma", "") + " " + x.get("izoh", "")).lower())]
    if not allr and o.saldo in (None, ""): raise SystemExit(f"{o.stir} bo'yicha bazada qator yo'q (kontragentlar buyrug'ini ko'ring)")
    prev = [x for x in allr if x["date"] < dan]; cur = sorted([x for x in allr if dan <= x["date"] <= gacha], key=lambda x: (x["date"], x["id"]))
    s0 = num(o.saldo) if o.saldo not in (None, "") else sum(x["summa"] if x["tur"] in DEBET else -x["summa"] for x in prev)
    knom = o.nom or db["nomlar"].get(o.stir) or f"STIR {o.stir}"
    if not o.nom and not re.search(r"(?i)\b(MCHJ|МЧЖ|ООО|OOO|AJ|АЖ|ATB|DM|XK|ХК|YATT|ЯТТ|QK|ЧП|ИП|MAS.ULIYATI)\b", knom):
        knom = f"ООО \"{knom.strip(chr(34))}\""
    TNR = "Times New Roman"; f = lambda b=False, s=10: Font(name=TNR, size=s, bold=b)
    th = Side(style="thin"); BD = Border(left=th, right=th, top=th, bottom=th)
    C = Alignment(horizontal="center", vertical="center", wrap_text=True); L = Alignment(horizontal="left", vertical="center", wrap_text=True)
    NUM = '#,##0.00;-#,##0.00;""'
    wb = Workbook(); ws = wb.active; ws.title = "Акт сверки"
    ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    # Sahifalash: fitToPage EMAS (u qo'lda bo'linishni buzadi) — qat'iy masshtab + o'zimiz hisoblagan bo'linishlar
    SCALE = 86
    ws.sheet_properties.pageSetUpPr.fitToPage = False; ws.page_setup.scale = SCALE
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.2, footer=0.2)
    ws.oddFooter.center.text = "Стр. &P из &N"; ws.oddFooter.center.size = 8
    ws.print_options.horizontalCentered = True
    for c, w in zip("ABCDEFGHI", [4, 9.5, 31, 14, 14, 9.5, 31, 14, 14]): ws.column_dimensions[c].width = w
    H = {}  # qator balandliklari (pt) — sahifalash uchun

    def m(rng, v, font=None, al=C, border=False, fill=None):
        ws.merge_cells(rng); c = ws[rng.split(":")[0]]; c.value = v; c.font = font or f(); c.alignment = al
        if fill: c.fill = fill
        if border:
            for row in ws[rng]:
                for x in row: x.border = BD
        return c
    m("A1:I1", "АКТ СВЕРКИ", f(True, 14)); H.update({1: 20, 2: 16, 3: 15, 5: 8, 7: 8, 8: 16, 9: 16})
    m("A2:I2", "взаимных расчетов", f(True, 11))
    m("A3:I3", f"за период: {dmy(dan)} - {dmy(gacha)}", f(False, 10))
    m("A4:I4", f"между {B['nom']} и {knom}" + (f"\nпо договору № {o.shartnoma}" if o.shartnoma else ""), f(True, 10)); H[4] = 30 if o.shartnoma else 16
    m("A6:I6", f"Мы, нижеподписавшиеся, {B['nom']}, с одной стороны, и {knom}, с другой стороны, составили настоящий акт сверки в том, "
               f"что состояние взаимных расчетов по данным учета следующее:", f(), L); H[6] = 30
    m("A8:E8", f"По данным {B['nom']}, сум", f(True), C, True); m("F8:I8", f"По данным {knom}, сум", f(True), C, True)
    # uzun nomli kontragent sarlavhasi kesilib qolmasin: ~55 belgi bir qatorga sig'adi
    H[8] = 16 if max(len(B['nom']), len(knom)) + 15 <= 55 else 30
    hdr = ["№", "Дата", "Документ", "Дебет", "Кредит", "Дата", "Документ", "Дебет", "Кредит"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(9, j, h); c.font = f(True); c.alignment = C; c.border = BD
    GR = PatternFill("solid", fgColor="F2F2F2")
    r = 10; H[10] = 16
    m(f"A{r}:C{r}", "Сальдо начальное", f(True), L, True, GR); m(f"F{r}:G{r}", "Сальдо начальное", f(True), L, True, GR)
    for j, v in ((4, s0 if s0 > 0 else None), (5, -s0 if s0 < 0 else None), (8, -s0 if s0 < 0 else None), (9, s0 if s0 > 0 else None)):
        c = ws.cell(r, j, v); c.number_format = NUM; c.font = f(True); c.border = BD; c.fill = GR
    first = r + 1; r += 1
    NAMES = {"sotdik": "Реализация", "oldik": "Поступление", "toladik": "Оплата", "toladi": "Поступление оплаты"}
    for i, x in enumerate(cur, 1):
        doc = f"{NAMES[x['tur']]} ({x.get('hujjat', '')} от {dmy(x['date'])})"
        deb = x["summa"] if x["tur"] in DEBET else None; kre = None if deb else x["summa"]
        vals = [i, dmy(x["date"]), doc, deb, kre, dmy(x["date"]), doc, kre, deb]
        H[r] = 26 if len(doc) > 31 else 15
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v); c.font = f(); c.border = BD
            c.alignment = L if j in (3, 7) else C
            if j in (4, 5, 8, 9): c.number_format = NUM; c.alignment = Alignment(horizontal="right", vertical="center")
        r += 1
    last = r - 1
    H[r] = 16; H[r + 1] = 16
    m(f"A{r}:C{r}", "Обороты за период", f(True), L, True, GR); m(f"F{r}:G{r}", "Обороты за период", f(True), L, True, GR)
    for j, L_ in ((4, "D"), (5, "E"), (8, "H"), (9, "I")):
        c = ws.cell(r, j, f"=SUM({L_}{first}:{L_}{last})" if last >= first else 0); c.number_format = NUM; c.font = f(True); c.border = BD; c.fill = GR
    ro = r; r += 1
    m(f"A{r}:C{r}", "Сальдо конечное", f(True), L, True, GR); m(f"F{r}:G{r}", "Сальдо конечное", f(True), L, True, GR)
    fs = f"(D10-E10+D{ro}-E{ro})"
    for j, frm in ((4, f"=MAX({fs},0)"), (5, f"=MAX(-{fs},0)"), (8, f"=MAX(-{fs},0)"), (9, f"=MAX({fs},0)")):
        c = ws.cell(r, j, frm); c.number_format = NUM; c.font = f(True); c.border = BD; c.fill = GR
    rs = r
    D = sum(x["summa"] for x in cur if x["tur"] in DEBET); K = sum(x["summa"] for x in cur if x["tur"] not in DEBET); s1 = s0 + D - K
    r += 2
    if abs(s1) < 0.005: txt = f"По данным {B['nom']} на {dmy(gacha)} задолженность отсутствует."
    elif s1 > 0: txt = f"По данным {B['nom']} на {dmy(gacha)} задолженность в пользу {B['nom']} составляет {money(s1)} сум ({propis(s1)})."
    else: txt = f"По данным {B['nom']} на {dmy(gacha)} задолженность в пользу {knom} составляет {money(-s1)} сум ({propis(s1)})."
    m(f"A{r}:I{r}", txt, f(True), L); H[r] = 30
    r += 3
    m(f"A{r}:E{r}", f"От {B['nom']}", f(True), L); m(f"F{r}:I{r}", f"От {knom}", f(True), L); r += 2
    m(f"A{r}:E{r}", f"Директор ________________ / {B.get('direktor', '')} /", f(), L); m(f"F{r}:I{r}", "Директор ________________ / ______________ /", f(), L); r += 2
    m(f"A{r}:E{r}", f"Гл. бухгалтер ____________ / {B.get('buxgalter', '') or '______________'} /", f(), L)
    m(f"F{r}:I{r}", "Гл. бухгалтер ____________ / ______________ /", f(), L); r += 2
    m(f"A{r}:E{r}", "М.П.", f(), L); m(f"F{r}:I{r}", "М.П.", f(), L)
    end = r
    ws.print_title_rows = "8:9"
    # --- sahifalash: jami+imzo bloki bo'linmaydi va undan oldin kamida 3 ta qator shu sahifada qoladi ---
    from openpyxl.worksheet.pagebreak import Break
    for rr_ in range(1, end + 1):
        ws.row_dimensions[rr_].height = H.get(rr_, 15)
    budget = (595.2 - 72 * (0.5 + 0.5)) / (SCALE / 100) - 12   # sahifa balandligi (varaq pt), zaxira bilan
    keep_from = max(first, last - 2) if last >= first else ro
    used = 0.0; breaks = []; pages = 1
    rr_ = 1
    while rr_ <= end:
        if rr_ == keep_from:
            blk = sum(ws.row_dimensions[k].height for k in range(rr_, end + 1))
            if used + blk > budget and rr_ > first:
                breaks.append(rr_ - 1); used = 32.0; pages += 1
            break
        h_ = ws.row_dimensions[rr_].height
        if used + h_ > budget and rr_ > 10:
            breaks.append(rr_ - 1); used = 32.0; pages += 1
        used += h_; rr_ += 1
    for b_ in breaks: ws.row_breaks.append(Break(id=b_))
    # Shartnomalar bo'yicha varaq (ma'lumot uchun)
    sj = shartnoma_jadval(cur)
    if sj:
        S2 = wb.create_sheet("По договорам")
        S2.page_setup.orientation = "landscape"; S2.page_setup.paperSize = S2.PAPERSIZE_A4
        S2.page_setup.fitToWidth = 1; S2.page_setup.fitToHeight = 0; S2.sheet_properties.pageSetUpPr.fitToPage = True
        S2.page_margins = PageMargins(left=0.6, right=0.6, top=0.6, bottom=0.6); S2.print_options.horizontalCentered = True
        NUM0 = '#,##0.00;-#,##0.00;0.00'
        S2.merge_cells("A1:E1"); c = S2["A1"]; c.value = f"Расшифровка по договорам (справочно): {B['nom']} — {knom}"; c.font = f(True, 11); c.alignment = C
        S2.merge_cells("A2:E2"); c = S2["A2"]; c.value = "Сальдо: «+» — задолженность контрагента перед нами, «−» — наша задолженность"; c.font = f(False, 9); c.alignment = C
        for j, h in enumerate(["Договор №", "Период", "Дебет", "Кредит", "Сальдо"], 1):
            c = S2.cell(4, j, h); c.font = f(True); c.border = BD; c.alignment = C
        rr = 5
        for k, (d_, k_, a, b) in sj.items():
            for j, v in enumerate([("без номера" if k == "raqamsiz" else k), f"{dmy(a)} – {dmy(b)}", d_, k_, round(d_ - k_, 2)], 1):
                c = S2.cell(rr, j, v); c.font = f(); c.border = BD; c.alignment = C if j <= 2 else Alignment(horizontal="right", vertical="center")
                if j >= 3: c.number_format = NUM0
            rr += 1
        for j, v in enumerate(["ИТОГО", "", f"=SUM(C5:C{rr - 1})", f"=SUM(D5:D{rr - 1})", f"=SUM(E5:E{rr - 1})"], 1):
            c = S2.cell(rr, j, v); c.font = f(True); c.border = BD; c.fill = GR
            if j >= 3: c.number_format = NUM0; c.alignment = Alignment(horizontal="right", vertical="center")
        for c_, w in zip("ABCDE", [18, 26, 20, 20, 20]): S2.column_dimensions[c_].width = w
    safe = re.sub(r"[^\w\-]+", "_", re.sub(r"ООО|OOO|MCHJ|\"", "", knom)).strip("_")[:40]
    out = os.path.join(HERE, f"AktSverka_{o.biz}_{safe}_{dmy(dan).replace('.', '')}-{dmy(gacha).replace('.', '')}.xlsx")
    wb.save(out)
    print(out)
    pdf = to_pdf(out)
    print(pdf if pdf else "PDF yasalmadi (Excel yoki LibreOffice topilmadi) — faqat Excel yuborildi")
    print(f"Sahifalar: akt {pages} bet" + (" + shartnomalar 1 bet" if sj else "") + " (imzo bloki oxirgi qatorlar bilan bitta betda)")
    print(f"Saldo boshi: {money(s0)} | Debet: {money(D)} | Kredit: {money(K)} | Saldo oxiri: {money(s1)} | qatorlar: {len(cur)}")
    print("Natija: " + ("qarz yo'q" if abs(s1) < 0.005 else (f"{knom} bizga ({B['nom']}) qarz: {money(s1)}" if s1 > 0 else f"Biz ({B['nom']}) {knom} ga qarz: {money(-s1)}")))
    bd = [x["date"] for x in db["qatorlar"] if x["biz"] == o.biz and x["manba"] == "bank"]
    if not bd: print("DIQQAT: bazada bu korxonaning bank ko'chirmasi yo'q — to'lovlar aktga kirmadi!")
    elif max(bd) < gacha: print(f"DIQQAT: bank ko'chirmasi faqat {dmy(max(bd))} gacha yuklangan — {dmy(max(bd))}–{dmy(gacha)} to'lovlari aktda yo'q. Yangi ko'chirmani yuklang.")
    elif min(bd) > dan: print(f"DIQQAT: bank ko'chirmasi {dmy(min(bd))} dan boshlanadi — undan oldingi to'lovlar yo'q.")
    fd = [x["date"] for x in db["qatorlar"] if x["biz"] == o.biz and x["manba"] == "faktura"]
    if not fd: print("DIQQAT: bazada fakturalar yo'q — faqat bank to'lovlari aktga kirdi!")
    if cur and not any(x["tur"] in ("toladik", "toladi") for x in cur):
        print("DIQQAT: aktda birorta ham TO'LOV yo'q — yuborishdan oldin 'tekshir' buyrug'ini ishga tushiring (bank ko'chirmasi yetishmasligi mumkin)")
    tash = [x for x in db["tashlangan"] if x["biz"] == o.biz and o.stir in (x.get("sel"), x.get("buy")) and x.get("date") and dan <= x["date"] <= gacha]
    if tash: print(f"Diqqat: shu kontragent bo'yicha {len(tash)} ta faktura tashlangan (tashlangan --stir {o.stir} bilan ko'ring)")
    if len(sj) > 1 or (sj and o.shartnoma == ""):
        print("Shartnomalar bo'yicha (+ = kontragent bizga qarz, − = biz qarzmiz):")
        for k, (d_, k_, a, b) in sj.items():
            print(f"  {k}: to'lov/sotuv {money(d_)} | faktura/tushum {money(k_)} | saldo {money(d_ - k_)}")


def to_pdf(xlsx):
    """Excel -> PDF: Windows'da Microsoft Excel orqali, bo'lmasa LibreOffice."""
    import subprocess, shutil
    pdf = os.path.splitext(xlsx)[0] + ".pdf"
    if os.name == "nt":
        ps = ("$ErrorActionPreference='Stop'; $x=New-Object -ComObject Excel.Application; $x.Visible=$false; $x.DisplayAlerts=$false; "
              f"$w=$x.Workbooks.Open('{xlsx}'); $w.ExportAsFixedFormat(0,'{pdf}'); $w.Close($false); $x.Quit()")
        try:
            subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True, capture_output=True, timeout=120)
            if os.path.exists(pdf): return pdf
        except Exception:
            pass
    so = shutil.which("soffice") or shutil.which("soffice.exe") or (r"C:\Program Files\LibreOffice\program\soffice.exe" if os.name == "nt" else None)
    if so and os.path.exists(so):
        subprocess.run([so, "--headless", "--convert-to", "pdf", "--outdir", os.path.dirname(xlsx), xlsx], capture_output=True, timeout=180)
        if os.path.exists(pdf): return pdf
    return None


def _nk(t):
    t = (t or "").upper()
    for w in ("MCHJ", "МЧЖ", "OOO", "ООО", "MAS'ULIYATI CHEKLANGAN JAMIYAT", "XK", "ХК", "AJ", "АЖ", "LLC"):
        t = t.replace(w, " ")
    return [w for w in re.findall(r"[A-ZА-ЯЁЎҚҒҲ0-9']{4,}", t)]


def cmd_tekshir(o, db):
    """Akt uchun mustaqil tekshiruv: aktga KIRMAGAN, lekin shu kontragentga tegishli bo'lishi mumkin bo'lgan qatorlarni topadi."""
    GEN = {"QURILISH", "SERVIS", "SERVICE", "TRANS", "INVEST", "BUILD", "STROY", "LOGISTIK", "LOGISTIC", "TEXNIK", "TECHNIC", "GROUP",
           "PLUS", "TRADE", "BETON", "TEMIR", "SAVDO", "AVTO", "AUTO", "TEXNIKA", "XIZMAT", "MCHJ", "КУРИЛИШ", "СЕРВИС", "ТРАНС", "COMPANY", "INTERNATIONAL"}
    st = o.stir; nom = db["nomlar"].get(st, ""); words = [w for w in _nk(nom) if len(w) >= 4 and w not in GEN][:3]
    mine = [x for x in db["qatorlar"] if x["biz"] == o.biz and x["stir"] == st]
    shs = {shartnoma_key(x) for x in mine} - {"raqamsiz"}
    muammo = 0
    print(f"TEKSHIRUV: {nom} ({st}) — bazada {len(mine)} qator; kalit so'zlar: {', '.join(words) or '-'}; shartnomalar: {', '.join(sorted(shs)) or '-'}")
    # 1) boshqa STIR (yoki 000000000) ga yozilgan, lekin nomi/maqsadida shu firma yoki uning shartnomasi bor qatorlar
    sus = []
    for x in db["qatorlar"]:
        if x["biz"] != o.biz or x["stir"] == st: continue
        t = ((x.get("nom") or "") + " " + db["nomlar"].get(x["stir"], "")).upper()
        hit_name = bool(words) and all(w in t for w in words)
        k = shartnoma_key(x)
        hit_sh = (not x["stir"].strip("0")) and k in shs and len(k) >= 4
        if hit_name or hit_sh: sus.append((x, "nomi" if hit_name else "shartnoma"))
    if sus:
        muammo += len(sus)
        print(f"1) DIQQAT: boshqa STIR ga yozilgan, lekin shu firmaga tegishli ko'rinadigan {len(sus)} qator (aktga KIRMAGAN):")
        for x, w in sus[:30]:
            print(f"   id {x['id']} | {dmy(x['date'])} | {x['tur']} | {money(x['summa'])} | STIR {x['stir']} | {w} mos | {(x.get('izoh') or '')[:90]}")
        print("   Tuzatish: haqiqatan shu firmaniki bo'lsa — python akt_tool.py kochir --id <id> --stir " + st)
    else:
        print("1) Boshqa STIR ga yozilgan shubhali qator yo'q ✓")
    # 2) bir xil sana+summa takrorlari
    g2 = collections.defaultdict(list)
    for x in mine: g2[(x["date"], x["tur"], round(x["summa"], 2))].append(x.get("hujjat", ""))
    dups = [(k, h) for k, h in g2.items() if len(h) > 1 and len(set(h)) < len(h)]
    same_day = [(k, h) for k, h in g2.items() if len(h) > 1 and len(set(h)) == len(h)]
    if dups:
        muammo += len(dups); print(f"2) DIQQAT: bir xil sana+summa+hujjat {len(dups)} marta (TAKROR):")
        for (d, t, sm), h in dups[:15]: print(f"   {dmy(d)} {t} {money(sm)} — {', '.join(h)}")
    else: print("2) Takror yo'q ✓" + (f" (bir kunda bir xil summa {len(same_day)} marta, lekin hujjat raqamlari har xil — normal)" if same_day else ""))
    # 3) tashlangan fakturalar
    tash = [x for x in db["tashlangan"] if x["biz"] == o.biz and st in (x.get("sel"), x.get("buy"))]
    nos = [x for x in tash if "SUMMA YOZILMAGAN" in x["sabab"]]
    print(f"3) Tashlangan fakturalar: {len(tash)}" + (f" — shundan {len(nos)} tasida SUMMA YO'Q (Juratdan so'rash kerak)" if nos else " ✓"))
    muammo += len(nos)
    # 4) bank va faktura qamrovi
    bd = [x["date"] for x in db["qatorlar"] if x["biz"] == o.biz and x["manba"] == "bank"]
    fd = [x for x in db["qatorlar"] if x["biz"] == o.biz and x["manba"] == "faktura"]
    sot = [x for x in fd if x["tur"] == "sotdik"]
    print(f"4) Bank: {dmy(min(bd)) + ' — ' + dmy(max(bd)) if bd else 'YO`Q'} | fakturalar: {len(fd)} (chiquvchi: {len(sot)})" + ("" if sot else " — DIQQAT: chiquvchi fakturalar yuklanmagan, biz sotgan xizmatlar aktda yo'q bo'lishi mumkin"))
    if not sot: muammo += 1
    # 5) bizning boshqa korxonalarimiz orqali to'lovlar
    F = firms(); others = [k for k in F if k != o.biz]
    oth = [x for x in db["qatorlar"] if x["biz"] in others and x["stir"] == st]
    if oth:
        print(f"5) DIQQAT: {', '.join(sorted({x['biz'] for x in oth}))} ham shu firma bilan {len(oth)} ta operatsiya qilgan — ularni aralashtirmang, lekin Juratga ayting")
    else: print("5) Boshqa korxonalarimiz bilan aralashgan operatsiya yo'q ✓")
    # 6) bir tomonlama akt: faqat faktura bor, to'lov yo'q (yoki aksincha) — odatda bank ko'chirmasi o'sha davr uchun yuklanmagan
    fk = [x for x in mine if x["tur"] in ("sotdik", "oldik")]; tl = [x for x in mine if x["tur"] in ("toladik", "toladi")]
    if fk and not tl:
        muammo += 1
        print(f"6) DIQQAT: {len(fk)} ta faktura bor, lekin BIRORTA HAM TO'LOV YO'Q. Bank ko'chirmasida shu firma nomi bilan qidiring "
              f"(nomi boshqacha yozilgan yoki STIR 000000000 bo'lishi mumkin), topilsa: kochir --id <id> --stir {st}")
    elif tl and not fk:
        muammo += 1; print(f"6) DIQQAT: {len(tl)} ta to'lov bor, lekin faktura yo'q — Didox eksportida shu firma fakturalarini tekshiring")
    else: print("6) Faktura ham, to'lov ham bor ✓")
    # 7) bank ko'chirmasi qamrovi kontragent bilan birinchi operatsiyani qoplaydimi
    if mine and bd and min(bd) > min(x["date"] for x in mine):
        muammo += 1
        print(f"7) DIQQAT: shu firma bilan birinchi operatsiya {dmy(min(x['date'] for x in mine))}, lekin bank ko'chirmasi faqat {dmy(min(bd))} dan yuklangan — "
              f"oldingi davr ko'chirmasini yuklang, aks holda o'sha to'lovlar aktga tushmaydi")
    else: print("7) Bank ko'chirmasi davri yetarli ✓")
    print(f"XULOSA: {'TO`G`RI ✓' if muammo == 0 else f'{muammo} ta tekshirilishi kerak bo`lgan joy'}")


def cmd_kochir(o, db):
    for x in db["qatorlar"]:
        if x["id"] == o.id:
            old = x["stir"]; x["stir"] = o.stir; save(db)
            print(f"id {o.id}: STIR {old} → {o.stir} ({dmy(x['date'])} {money(x['summa'])})"); return
    print("topilmadi")


def cmd_tashlangan(o, db):
    for x in db["tashlangan"]:
        if x["biz"] == o.biz and (not o.stir or o.stir in (x.get("sel"), x.get("buy"))):
            print(f"СФ №{x.get('no')} {dmy(x['date']) if x.get('date') else ''} {money(x.get('sum') or 0)} — {x['sabab']}")


def cmd_holat(o, db):
    c = collections.Counter((x["biz"], x["manba"]) for x in db["qatorlar"])
    for (b, mn), n in sorted(c.items()):
        ds = [x["date"] for x in db["qatorlar"] if x["biz"] == b and x["manba"] == mn]
        print(f"{b} | {mn}: {n} qator | {dmy(min(ds))} — {dmy(max(ds))}")
    if not c: print("Baza bo'sh — bank ko'chirmasi va fakturalarni yuklang.")


def main():
    ap = argparse.ArgumentParser(); sp = ap.add_subparsers(dest="cmd", required=True)
    for n in ("bank", "faktura"):
        p = sp.add_parser(n); p.add_argument("--biz", required=True); p.add_argument("--file", required=True)
    q = sp.add_parser("qator")
    for n in ("--biz", "--stir", "--date", "--tur", "--summa"): q.add_argument(n, required=True)
    for n in ("--hujjat", "--shartnoma", "--nom"): q.add_argument(n, default="")
    k = sp.add_parser("kontragentlar"); k.add_argument("--biz", required=True); k.add_argument("--qidir", default="")
    a = sp.add_parser("akt")
    for n in ("--biz", "--stir"): a.add_argument(n, required=True)
    for n in ("--dan", "--gacha", "--saldo", "--shartnoma", "--nom", "--sana"): a.add_argument(n, default="")
    t = sp.add_parser("tashlangan"); t.add_argument("--biz", required=True); t.add_argument("--stir", default="")
    d = sp.add_parser("del"); d.add_argument("--id", type=int, required=True)
    sp.add_parser("holat"); sp.add_parser("inn_tikla")
    tk = sp.add_parser("tekshir"); tk.add_argument("--biz", required=True); tk.add_argument("--stir", required=True)
    kc = sp.add_parser("kochir"); kc.add_argument("--id", type=int, required=True); kc.add_argument("--stir", required=True)
    tz = sp.add_parser("tozala"); tz.add_argument("--biz", required=True); tz.add_argument("--manba", required=True, choices=["bank", "faktura", "qo'lda"])
    o = ap.parse_args(); db = load()
    if hasattr(o, "biz") and o.biz:
        o.biz = o.biz.upper()
        if o.biz not in firms(): raise SystemExit("--biz: " + " | ".join(firms()))
    if o.cmd == "inn_tikla":
        acc2stir = db.get("hisoblar", {}); n = 0
        nom2stir = {re.sub(r"[^A-ZА-Я0-9]", "", v.upper()): k for k, v in db["nomlar"].items() if k.strip("0")}
        for x in db["qatorlar"]:
            if x["stir"].strip("0"): continue
            nk = re.sub(r"[^A-ZА-Я0-9]", "", (x.get("nom") or "").upper())
            real = acc2stir.get(x.get("hisob", "")) or nom2stir.get(nk)
            if real: x["stir"] = real; n += 1
        save(db); print(f"INN tiklandi: {n} qator"); return
    if o.cmd == "tozala":
        n = len(db["qatorlar"]); db["qatorlar"] = [x for x in db["qatorlar"] if not (x["biz"] == o.biz and x["manba"] == o.manba)]
        if o.manba == "faktura": db["tashlangan"] = [x for x in db["tashlangan"] if x["biz"] != o.biz]
        save(db); print(f"{o.biz} {o.manba}: {n - len(db['qatorlar'])} qator o'chirildi (qayta yuklang)"); return
    if o.cmd == "del":
        n = len(db["qatorlar"]); db["qatorlar"] = [x for x in db["qatorlar"] if x["id"] != o.id]; save(db)
        print("o'chirildi" if len(db["qatorlar"]) < n else "topilmadi"); return
    {"bank": cmd_bank, "faktura": cmd_faktura, "qator": cmd_qator, "kontragentlar": cmd_kontragentlar, "akt": cmd_akt,
     "tashlangan": cmd_tashlangan, "holat": cmd_holat, "tekshir": cmd_tekshir, "kochir": cmd_kochir}[o.cmd](o, db)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"): sys.stdout.reconfigure(encoding="utf-8")
    main()
